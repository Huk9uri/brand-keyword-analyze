from __future__ import annotations

import hashlib
import math
import re
from abc import ABC, abstractmethod
from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from pathlib import Path
from types import MappingProxyType
from typing import Any

from openpyxl import load_workbook

from .matching import OwnStore
from .models import DeviceType, SearchTask, SourceType


EXCEL_HEADERS = (
    "검색어",
    "기준매장ID",
    "기준매장명",
    "조회위치명",
    "위도",
    "경도",
    "기기",
    "소스ID",
)
REQUIRED_EXCEL_HEADERS = EXCEL_HEADERS[:4]
OWN_STORE_SHEET = "자사매장"
OWN_STORE_HEADERS = ("매장ID", "매장명", "네이버플레이스ID", "수동매핑플레이스ID")


@dataclass(frozen=True)
class ExcelWorkbookInput:
    tasks: tuple[SearchTask, ...]
    own_stores: tuple[OwnStore, ...]
    manual_place_id_mappings: Mapping[str, str]

    def __post_init__(self) -> None:
        object.__setattr__(
            self,
            "manual_place_id_mappings",
            MappingProxyType(dict(self.manual_place_id_mappings)),
        )


class InputValidationError(ValueError):
    """Raised when one or more input rows cannot become search tasks."""

    def __init__(self, errors: Sequence[str]) -> None:
        self.errors = tuple(errors)
        super().__init__("input validation failed:\n" + "\n".join(self.errors))


class TaskSource(ABC):
    @abstractmethod
    def load(self) -> list[SearchTask]:
        """Load, validate, and deduplicate search tasks."""


class ExcelTaskSource(TaskSource):
    def __init__(self, path: str | Path) -> None:
        self.path = Path(path)

    def load(self) -> list[SearchTask]:
        workbook = load_workbook(self.path, read_only=True, data_only=True)
        try:
            if "검색작업" not in workbook.sheetnames:
                raise InputValidationError(["required worksheet is missing: 검색작업"])
            worksheet = workbook["검색작업"]
            rows = worksheet.iter_rows(values_only=True)
            try:
                header_row = next(rows)
            except StopIteration as exc:
                raise InputValidationError(["row 1: header row is required"]) from exc

            headers = {str(value).strip(): index for index, value in enumerate(header_row) if value is not None}
            missing_headers = [header for header in REQUIRED_EXCEL_HEADERS if header not in headers]
            if missing_headers:
                raise InputValidationError(["row 1: missing required headers: " + ", ".join(missing_headers)])

            tasks: list[SearchTask] = []
            errors: list[str] = []
            seen_task_keys: set[str] = set()
            for row_number, row in enumerate(rows, start=2):
                values = {header: _value_at(row, headers.get(header)) for header in EXCEL_HEADERS}
                if not any(value is not None and str(value).strip() for value in values.values()):
                    continue
                source_id = _text(values["소스ID"]) or f"excel-row-{row_number}"
                task, row_errors = _build_task(
                    values={
                        "query": values["검색어"],
                        "target_store_id": values["기준매장ID"],
                        "target_store_name": values["기준매장명"],
                        "location_name": values["조회위치명"],
                        "latitude": values["위도"],
                        "longitude": values["경도"],
                        "device_type": values["기기"],
                    },
                    task_id=source_id,
                    source_type=SourceType.EXCEL,
                    source_id=source_id,
                )
                if row_errors:
                    errors.extend(f"row {row_number}: {error}" for error in row_errors)
                    continue
                assert task is not None
                task_key = build_task_key(task)
                if task_key not in seen_task_keys:
                    seen_task_keys.add(task_key)
                    tasks.append(task)

            if errors:
                raise InputValidationError(errors)
            return tasks
        finally:
            workbook.close()

    def load_bundle(self) -> ExcelWorkbookInput:
        """Load task rows plus the separate authoritative own-store identity sheet."""

        tasks = tuple(self.load())
        workbook = load_workbook(self.path, read_only=True, data_only=True)
        try:
            if OWN_STORE_SHEET not in workbook.sheetnames:
                raise InputValidationError([f"required worksheet is missing: {OWN_STORE_SHEET}"])
            own_stores, manual_mappings = _load_own_store_identity(workbook[OWN_STORE_SHEET])
        finally:
            workbook.close()

        missing_store_ids = sorted(
            {task.target_store_id for task in tasks} - {store.store_id for store in own_stores}
        )
        if missing_store_ids:
            raise InputValidationError(
                [
                    f"{OWN_STORE_SHEET}: missing identity for task store IDs: "
                    + ", ".join(missing_store_ids)
                ]
            )
        return ExcelWorkbookInput(tasks, own_stores, manual_mappings)


class DatabaseTaskSource(TaskSource):
    """Adapt database candidate records without exposing their column names downstream."""

    def __init__(self, cursor: Any, query: str) -> None:
        self.cursor = cursor
        self.query = query

    def load(self) -> list[SearchTask]:
        self.cursor.execute(self.query)
        rows = self.cursor.fetchall()
        columns = _cursor_columns(self.cursor)
        tasks: list[SearchTask] = []
        errors: list[str] = []
        seen_task_keys: set[str] = set()
        for row_number, row in enumerate(rows, start=1):
            record = _row_mapping(row, columns)
            source_id = _text(record.get("candidate_id")) or f"database-row-{row_number}"
            task, row_errors = _build_task(
                values={
                    "query": record.get("candidate_keyword"),
                    "target_store_id": record.get("store_code"),
                    "target_store_name": record.get("store_display_name"),
                    "location_name": record.get("search_location"),
                    "latitude": record.get("search_latitude"),
                    "longitude": record.get("search_longitude"),
                    "device_type": record.get("device_code"),
                },
                task_id=f"database-{source_id}",
                source_type=SourceType.DATABASE,
                source_id=source_id,
            )
            if row_errors:
                errors.extend(f"database row {row_number}: {error}" for error in row_errors)
                continue
            assert task is not None
            task_key = build_task_key(task)
            if task_key not in seen_task_keys:
                seen_task_keys.add(task_key)
                tasks.append(task)

        if errors:
            raise InputValidationError(errors)
        return tasks


def normalize_query(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip()


def _load_own_store_identity(worksheet: Any) -> tuple[tuple[OwnStore, ...], dict[str, str]]:
    rows = worksheet.iter_rows(values_only=True)
    try:
        header_row = next(rows)
    except StopIteration as exc:
        raise InputValidationError([f"{OWN_STORE_SHEET} row 1: header row is required"]) from exc

    headers = {str(value).strip(): index for index, value in enumerate(header_row) if value is not None}
    missing_headers = [header for header in OWN_STORE_HEADERS if header not in headers]
    if missing_headers:
        raise InputValidationError(
            [f"{OWN_STORE_SHEET} row 1: missing required headers: " + ", ".join(missing_headers)]
        )

    stores_by_id: dict[str, OwnStore] = {}
    canonical_place_ids: dict[str, str] = {}
    manual_mappings: dict[str, str] = {}
    errors: list[str] = []
    for row_number, row in enumerate(rows, start=2):
        values = {
            header: _text(_value_at(row, headers[header]))
            for header in OWN_STORE_HEADERS
        }
        if not any(values.values()):
            continue
        store_id = values["매장ID"]
        store_name = values["매장명"]
        place_id = values["네이버플레이스ID"]
        manual_place_id = values["수동매핑플레이스ID"]
        missing = [
            name
            for value, name in (
                (store_id, "store_id"),
                (store_name, "store_name"),
                (place_id, "naver_place_id"),
            )
            if not value
        ]
        if missing:
            errors.append(
                f"{OWN_STORE_SHEET} row {row_number}: required values missing: "
                + ", ".join(missing)
            )
            continue

        store = OwnStore(store_id, store_name, place_id)
        existing_store = stores_by_id.get(store_id)
        if existing_store is not None and existing_store != store:
            errors.append(f"{OWN_STORE_SHEET} row {row_number}: conflicting identity for {store_id}")
            continue
        other_store_id = canonical_place_ids.get(place_id)
        if other_store_id is not None and other_store_id != store_id:
            errors.append(
                f"{OWN_STORE_SHEET} row {row_number}: place ID {place_id} belongs to multiple stores"
            )
            continue
        stores_by_id.setdefault(store_id, store)
        canonical_place_ids[place_id] = store_id

        if manual_place_id:
            mapped_store_id = manual_mappings.get(manual_place_id)
            if mapped_store_id is not None and mapped_store_id != store_id:
                errors.append(
                    f"{OWN_STORE_SHEET} row {row_number}: manual place ID {manual_place_id} "
                    "belongs to multiple stores"
                )
                continue
            manual_mappings[manual_place_id] = store_id

    if errors:
        raise InputValidationError(errors)
    if not stores_by_id:
        raise InputValidationError([f"{OWN_STORE_SHEET}: at least one own store is required"])
    return tuple(stores_by_id.values()), manual_mappings


def build_task_key(task: SearchTask) -> str:
    latitude = "AUTO" if task.latitude is None else f"{task.latitude:.7f}"
    longitude = "AUTO" if task.longitude is None else f"{task.longitude:.7f}"
    key = "|".join(
        (
            normalize_query(task.query),
            task.target_store_id,
            latitude,
            longitude,
            task.device_type.value,
        )
    )
    return hashlib.sha256(key.encode("utf-8")).hexdigest()


def _build_task(
    values: Mapping[str, object],
    *,
    task_id: str,
    source_type: SourceType,
    source_id: str,
) -> tuple[SearchTask | None, list[str]]:
    query = normalize_query(_text(values.get("query")))
    target_store_id = _text(values.get("target_store_id"))
    target_store_name = _text(values.get("target_store_name"))
    location_name = _text(values.get("location_name"))
    raw_latitude = values.get("latitude")
    raw_longitude = values.get("longitude")
    latitude_blank = raw_latitude is None or _text(raw_latitude) == ""
    longitude_blank = raw_longitude is None or _text(raw_longitude) == ""
    coordinate_pair_error = None
    if latitude_blank != longitude_blank:
        latitude = longitude = None
        latitude_error = longitude_error = None
        coordinate_pair_error = "latitude and longitude must both be provided or both be blank"
    elif latitude_blank:
        latitude = longitude = None
        latitude_error = longitude_error = None
    else:
        latitude, latitude_error = _coordinate(raw_latitude, "latitude", -90, 90)
        longitude, longitude_error = _coordinate(raw_longitude, "longitude", -180, 180)
    device_type, device_error = _device_type(values.get("device_type"))

    errors = []
    for value, name in (
        (query, "query"),
        (target_store_id, "target_store_id"),
        (target_store_name, "target_store_name"),
        (location_name, "location_name"),
    ):
        if not value:
            errors.append(f"{name} is required")
    if latitude_error:
        errors.append(latitude_error)
    if longitude_error:
        errors.append(longitude_error)
    if coordinate_pair_error:
        errors.append(coordinate_pair_error)
    if device_error:
        errors.append(device_error)
    if errors:
        return None, errors

    assert device_type is not None
    return (
        SearchTask(
            task_id=task_id,
            query=query,
            target_store_id=target_store_id,
            target_store_name=target_store_name,
            location_name=location_name,
            latitude=latitude,
            longitude=longitude,
            source_type=source_type,
            source_id=source_id,
            device_type=device_type,
        ),
        [],
    )


def _coordinate(value: object, name: str, lower: float, upper: float) -> tuple[float | None, str | None]:
    try:
        if isinstance(value, bool) or value is None or _text(value) == "":
            raise ValueError
        coordinate = float(value)
    except (TypeError, ValueError):
        return None, f"{name} must be a number"
    if not math.isfinite(coordinate) or not lower <= coordinate <= upper:
        return None, f"{name} must be between {lower:g} and {upper:g}"
    return coordinate, None


def _device_type(value: object) -> tuple[DeviceType | None, str | None]:
    normalized = _text(value).upper()
    if not normalized:
        return DeviceType.MOBILE, None
    try:
        return DeviceType(normalized), None
    except ValueError:
        return None, "device must be MOBILE or PC"


def _text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _value_at(row: tuple[object, ...], index: int | None) -> object | None:
    return row[index] if index is not None and index < len(row) else None


def _cursor_columns(cursor: Any) -> tuple[str, ...] | None:
    description = getattr(cursor, "description", None)
    if not description:
        return None
    return tuple(column[0] for column in description)


def _row_mapping(row: object, columns: tuple[str, ...] | None) -> Mapping[str, object]:
    if isinstance(row, Mapping):
        return row
    if columns is None:
        raise TypeError("database cursor rows must be mappings or expose cursor.description")
    return dict(zip(columns, row, strict=True))
